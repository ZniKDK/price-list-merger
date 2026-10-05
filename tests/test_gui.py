"""Тесты окна программы без экрана (QT_QPA_PLATFORM=offscreen).

Диалоги (сообщения, выбор файла) подменяются, чтобы тест не ждал нажатия кнопки.
"""

import os
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pandas as pd  # noqa: E402
import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PyQt6.QtCore import QSettings, Qt  # noqa: E402
from PyQt6.QtWidgets import QApplication, QFileDialog, QMessageBox  # noqa: E402

from price_list_merger.core.mapping import PRICE, SKU  # noqa: E402
from price_list_merger.core.pricing import ROUND_TEN  # noqa: E402
from price_list_merger.gui.main_window import (  # noqa: E402
    COLUMNS, FILES, RESULT, SETTINGS, MainWindow,
)
from price_list_merger.gui.style import apply_theme  # noqa: E402
from price_list_merger.gui.table_model import (  # noqa: E402
    DataFrameModel, SearchProxyModel, format_value,
)
from price_list_merger.profiles import ProfileStore  # noqa: E402


@pytest.fixture(scope="session")
def app():
    from price_list_merger.gui import animations

    animations.set_enabled(False)  # в тестах сразу конечное состояние, без ожидания
    application = QApplication.instance() or QApplication([])
    apply_theme(application, "light")
    return application


@pytest.fixture
def messages(monkeypatch):
    """Подменяет окна сообщений: вместо показа запоминает текст."""
    shown = []

    def fake(parent, title, text, *args, **kwargs):
        shown.append((title, text))
        return QMessageBox.StandardButton.Ok

    for name in ("warning", "critical", "information"):
        monkeypatch.setattr(QMessageBox, name, staticmethod(fake))
    return shown


@pytest.fixture
def window(app, tmp_path, messages):
    settings = QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat)
    win = MainWindow(ProfileStore(tmp_path / "profiles.json"), settings)
    yield win
    win.close()


def wait_merge(app, window, timeout=20):
    deadline = time.monotonic() + timeout
    while window.worker is not None:
        assert time.monotonic() < deadline, "свод не закончился вовремя"
        app.processEvents()
        time.sleep(0.01)
    app.processEvents()


def test_start_state(window):
    assert window.pages.currentIndex() == FILES
    assert not window.next_button.isEnabled()
    window.go_to(COLUMNS)  # без прайсов дальше не пускает
    assert window.pages.currentIndex() == FILES


def test_add_files_and_rename(window, demo_files, messages):
    page = window.files_page
    page.add_files(demo_files)
    page.add_files(demo_files[:1])  # повторно тот же файл не добавляется
    assert [s.supplier for s in window.sources] == ["ЭлектроОпт", "СтройСнаб", "ТехноМаркет"]
    assert page.table.rowCount() == 3
    assert window.next_button.isEnabled()

    page.table.item(1, 0).setText("Снаб")
    assert window.sources[1].supplier == "Снаб"
    page.table.item(1, 0).setText("ЭлектроОпт")  # такое имя уже есть
    assert window.sources[1].supplier == "Снаб"
    assert messages and messages[-1][0] == "Имя поставщика"

    page.table.selectRow(2)
    page.move_selected(-1)
    assert window.sources[1].supplier == "ТехноМаркет"
    page.remove_selected()
    assert len(window.sources) == 2


def test_bad_file_reported(window, tmp_path, messages):
    bad = tmp_path / "плохой.xlsx"
    bad.write_bytes(b"not excel")
    window.files_page.add_files([bad])
    assert window.sources == []
    assert "Не удалось открыть" in messages[-1][1]


def test_columns_page_changes_mapping(window, demo_files, messages):
    window.files_page.add_files(demo_files)
    window.go_to(COLUMNS)
    page = window.columns_page
    assert page.supplier_box.count() == 3
    assert page.header_spin.value() == 5  # у ЭлектроОпт шапка на 5-й строке
    page.field_boxes[PRICE].setCurrentIndex(0)  # «не используется»
    assert PRICE not in window.sources[0].mapping
    assert "цен" in page.status.text()

    window.go_to(SETTINGS)
    window.run_merge()  # с ошибкой в колонках свод не запускается
    assert window.pages.currentIndex() == COLUMNS
    assert "Не все колонки" in messages[-1][0]

    page.field_boxes[PRICE].setCurrentText("Цена, руб.")
    assert window.sources[0].mapping[PRICE] == "Цена, руб."
    page.header_spin.setValue(1)  # неверная строка заголовка — колонки подбираются заново
    assert window.sources[0].header_row == 0


def test_full_merge_and_save(app, window, demo_files, tmp_path, monkeypatch):
    window.files_page.add_files(demo_files)
    window.go_to(SETTINGS)
    settings = window.settings_page
    settings.percent.setValue(10)
    settings.rounding.setCurrentIndex(settings.rounding.findData(ROUND_TEN))
    window.run_merge()
    wait_merge(app, window)

    assert window.pages.currentIndex() == RESULT
    report = window.report
    assert report is not None
    result = window.result_page
    assert result.products_model.rowCount() == len(report.products)
    assert result.offers_model.rowCount() == len(report.offers)
    assert result.cards["products"].value.text() == str(len(report.products))
    # Профили сохранены — в следующий раз колонки подставятся сами
    assert window.profiles.get("СтройСнаб").mapping[SKU] == "Код товара"

    # Поиск фильтрует таблицу
    result.search.setText("Розетка")
    assert 0 < result.proxies[0].rowCount() < len(report.products)

    target = tmp_path / "итог.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(target), "")))
    result.save_button.click()
    assert target.exists()
    assert result.open_button.isVisibleTo(result)
    sheet = load_workbook(target)["Итог"]
    assert all(row[8].value % 10 == 0 for row in sheet.iter_rows(min_row=2) if row[8].value)


def test_settings_roundtrip(app, tmp_path):
    from price_list_merger.gui.settings_page import SettingsPage

    store = QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat)
    page = SettingsPage(store)
    page.percent.setValue(33)
    page.fuzzy.setChecked(False)
    page._add_tier(5000, 15)
    page.save()
    again = SettingsPage(store)
    assert again.percent.value() == 33
    assert not again.merge_settings().use_fuzzy
    assert (5000, 15) in [(t.min_price, t.percent) for t in again.pricing().tiers]
    assert "→" in again.example.text()


def test_format_value():
    assert format_value(1234.5, money=True) == "1 234,50"
    assert format_value(float("nan")) == ""
    assert format_value(None) == ""
    assert format_value(True) == "да"
    assert format_value(3.0) == "3"


def test_proxy_sorts_numbers_not_text(app):
    model = DataFrameModel(pd.DataFrame({"price": [100.0, 9.0, None, 25.0]}))
    proxy = SearchProxyModel()
    proxy.setSourceModel(model)
    proxy.sort(0, Qt.SortOrder.AscendingOrder)
    values = [proxy.index(row, 0).data() for row in range(proxy.rowCount())]
    assert values[:3] == ["9", "25", "100"]


def test_files_from_command_line(app, demo_files, tmp_path, monkeypatch):
    """Файлы, перетащенные на значок программы, сразу попадают в список."""
    import sys

    from price_list_merger import __main__ as entry
    from price_list_merger.gui import main_window

    opened = []
    real_init = main_window.MainWindow.__init__

    def init(self, *args, **kwargs):
        # Профили и настройки — во временной папке, а не в настоящих
        store = ProfileStore(tmp_path / "profiles.json")
        settings = QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat)
        real_init(self, store, settings)
        opened.append(self)

    monkeypatch.setattr(main_window.MainWindow, "__init__", init)
    import PyQt6.QtWidgets as widgets

    # Второе приложение Qt создать нельзя — main() получит уже готовое,
    # а цикл событий (exec) сразу завершится
    monkeypatch.setattr(widgets, "QApplication", lambda argv: app)
    monkeypatch.setattr(QApplication, "exec", lambda self: 0)
    monkeypatch.setattr(sys, "argv", ["prog", str(demo_files[1]), "нет-такого-файла.xlsx"])
    assert entry.main() == 0
    assert [source.supplier for source in opened[0].sources] == ["СтройСнаб"]
    opened[0].close()


def test_theme_switch(app, window):
    from PyQt6.QtGui import QPalette

    from price_list_merger.gui.style import DARK, LIGHT, C

    window.set_theme("dark")
    assert C["bg"] == DARK["bg"]
    assert app.palette().color(QPalette.ColorRole.Window).name().upper() == DARK["bg"].upper()
    assert window.settings.value("ui/theme") == "dark"
    assert window.theme_actions["dark"].isChecked()
    assert "тёмная" in window.sidebar.theme_button.text()

    window.set_theme("light")
    assert C["bg"] == LIGHT["bg"]
    assert window.settings.value("ui/theme") == "light"

    window.set_theme("system")  # тема Windows: светлая или тёмная — главное, без ошибок
    assert C["bg"] in (LIGHT["bg"], DARK["bg"])
    window.set_theme("light")


def test_highlight_follows_theme(app, window):
    """Подсветка ячеек хранится как смысл и перекрашивается при смене темы."""
    from price_list_merger.gui.style import DARK, LIGHT

    model = DataFrameModel(pd.DataFrame({"price": [1.0]}))
    model.set_cell_color(0, 0, "best")
    index = model.index(0, 0)
    window.set_theme("light")
    assert model.data(index, Qt.ItemDataRole.BackgroundRole).color().name().upper() == LIGHT["best_bg"].upper()
    window.set_theme("dark")
    assert model.data(index, Qt.ItemDataRole.BackgroundRole).color().name().upper() == DARK["best_bg"].upper()
    assert model.data(index, Qt.ItemDataRole.ForegroundRole).color().name().upper() == DARK["best_text"].upper()
    window.set_theme("light")


def test_no_top_menu_and_sidebar_tools(window):
    # Верхнее меню выбивалось из оформления — всё перенесено в боковую панель
    assert window.menuBar().actions() == []
    assert window.sidebar.profiles_button.isVisible() or not window.isVisible()
    assert [b.isEnabled() for b in window.sidebar.step_buttons] == [True, False, False, False]


def test_toast_after_save(app, window, demo_files, tmp_path, monkeypatch):
    window.files_page.add_files(demo_files)
    window.go_to(SETTINGS)
    window.run_merge()
    wait_merge(app, window)
    target = tmp_path / "свод.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(target), "")))
    window.save_excel()
    assert window.toast.isVisibleTo(window)
    assert "свод.xlsx" in window.toast.text.text()


def test_empty_hints(app, window, demo_files):
    result = window.result_page
    assert result.views[0].empty_hint[0] == "Свод ещё не выполнен"
    window.files_page.add_files(demo_files)
    window.go_to(SETTINGS)
    window.run_merge()
    wait_merge(app, window)
    result.search.setText("такого товара нет")
    assert result.proxies[0].rowCount() == 0
    assert result.views[0].empty_hint[0] == "Ничего не найдено"
    result.search.clear()
    assert result.views[0].empty_hint[0] == "Свод ещё не выполнен" or result.proxies[0].rowCount() > 0


def test_fit_columns_not_narrower_than_header(app):
    from PyQt6.QtGui import QFont, QFontMetrics

    from price_list_merger.gui.widgets import HintTableView, fit_columns

    model = DataFrameModel(pd.DataFrame({"offers": [1, 2]}), headers={"offers": "Поставщиков"})
    view = HintTableView()
    view.setModel(model)
    fit_columns(view)
    bold = QFont(view.font())
    bold.setWeight(QFont.Weight.DemiBold)
    assert view.columnWidth(0) >= QFontMetrics(bold).horizontalAdvance("Поставщиков") + 16
