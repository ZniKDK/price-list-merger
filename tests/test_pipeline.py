"""Тесты всего процесса и выгрузки в Excel на демо-прайсах."""

import pytest
from openpyxl import load_workbook

from price_list_merger.core.exporter import BEST_FILL, export_report
from price_list_merger.core.loader import PriceListError
from price_list_merger.core.mapping import PRICE
from price_list_merger.core.merging import MergeSettings
from price_list_merger.core.pipeline import SupplierSource, run_merge
from price_list_merger.core.pricing import ROUND_NINE, MarkupTier, PricingSettings


@pytest.fixture
def sources(demo_files):
    return [SupplierSource.from_file(path) for path in demo_files]


def test_demo_files_detected(sources):
    by_name = {source.supplier: source for source in sources}
    assert set(by_name) == {"ЭлектроОпт", "СтройСнаб", "ТехноМаркет"}
    assert by_name["ЭлектроОпт"].header_row == 4
    assert by_name["СтройСнаб"].table.raw.encoding == "cp1251"
    assert "sku" not in by_name["ТехноМаркет"].mapping
    assert all(not source.problems() for source in sources)


def test_demo_merge(sources):
    pricing = PricingSettings(percent=20, tiers=[MarkupTier(1000, 12)], rounding=ROUND_NINE)
    progress = []
    report = run_merge(sources, pricing=pricing, progress=lambda share, text: progress.append(share))
    stats = report.stats()
    assert stats["suppliers"] == 3
    assert stats["products"] < stats["offers"]  # товары действительно сведены
    assert stats["in_several"] >= 20
    assert stats["fuzzy"] >= 3
    assert stats["rejected"] >= 4
    assert progress[-1] == 1.0
    # Лучшая цена — минимум из цен поставщиков, цена продажи с наценкой и на …9
    for _, product in report.products.dropna(subset=["best_price"]).iterrows():
        prices = [product[f"price:{name}"] for name in report.suppliers]
        assert product["best_price"] == min(p for p in prices if p == p)
        assert product["sale_price"] >= product["best_price"]
        assert int(product["sale_price"]) % 10 == 9


def test_cyrillic_lookalike_sku_matched(sources):
    report = run_merge(sources)
    product = report.products[report.products["name"] == "Кабель-канал 16х16 мм"]
    assert len(product) == 1
    assert product.iloc[0]["offers"] == 3


def test_merge_without_fuzzy_gives_more_products(sources):
    strict = run_merge(sources, MergeSettings(use_fuzzy=False)).stats()["products"]
    fuzzy = run_merge(sources).stats()["products"]
    assert strict > fuzzy


def test_validation_errors(sources):
    with pytest.raises(PriceListError, match="ни одного"):
        run_merge([])
    sources[1].supplier = sources[0].supplier
    with pytest.raises(PriceListError, match="повторяются"):
        run_merge(sources)
    sources[1].supplier = "Другой"
    sources[0].mapping.pop(PRICE)
    with pytest.raises(PriceListError, match="цен"):
        run_merge(sources)


def test_reload_drops_missing_columns(sources):
    source = sources[0]
    source.header_row = 0  # заведомо неверная строка заголовка
    source.reload()
    assert all(column in source.table.columns for column in source.mapping.values())


def test_export(sources, tmp_path):
    report = run_merge(sources, pricing=PricingSettings(percent=10))
    path = export_report(report, tmp_path / "out" / "итог")
    assert path.suffix == ".xlsx"
    book = load_workbook(path)
    assert book.sheetnames == ["Итог", "Все предложения", "Не сопоставлено", "Отброшено", "Сводка"]

    sheet = book["Итог"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:9] == [
        "№", "Артикул", "Наименование", "Ед.", "Поставщиков",
        "Лучшая цена", "Поставщик", "Наценка, %", "Цена продажи",
    ]
    assert sheet.max_row == len(report.products) + 1
    assert sheet.freeze_panes == "A2"
    # В каждой строке с лучшей ценой подсвечена ячейка этого поставщика
    for row in sheet.iter_rows(min_row=2):
        supplier = row[6].value
        if supplier:
            cell = row[9 + report.suppliers.index(supplier)]
            assert cell.fill.fgColor.rgb.endswith(BEST_FILL.fgColor.rgb[-6:])
            assert cell.value == row[5].value

    assert book["Все предложения"].max_row == len(report.offers) + 1
    assert book["Отброшено"].max_row == len(report.rejected) + 1
    unmatched = (report.products["offers"] == 1).sum()
    assert book["Не сопоставлено"].max_row == unmatched + 1
    summary = {row[0].value: row[1].value for row in book["Сводка"].iter_rows(min_row=2)}
    assert summary["Поставщиков"] == 3
    assert summary["Наценка по умолчанию, %"] == 10
