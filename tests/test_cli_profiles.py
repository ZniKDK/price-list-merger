"""Тесты профилей поставщиков и командной строки."""

import json

import pytest
from openpyxl import load_workbook

from price_list_merger import cli
from price_list_merger.core.mapping import NAME, PRICE, SKU
from price_list_merger.core.pipeline import SupplierSource
from price_list_merger.profiles import ProfileStore, SupplierProfile


@pytest.fixture
def store_path(tmp_path, monkeypatch):
    """Профили — во временной папке, а не в настоящем %APPDATA%."""
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    return tmp_path / "appdata" / "PriceListMerger" / "profiles.json"


def test_profile_store_roundtrip(store_path):
    store = ProfileStore()
    assert store.path == store_path
    store.save("ЭлектроОпт", SupplierProfile({SKU: "Арт", PRICE: "Цена"}, header_row=3, sheet="Прайс"))
    again = ProfileStore()
    profile = again.get("  электроопт ")  # без учёта регистра и пробелов
    assert profile == SupplierProfile({SKU: "Арт", PRICE: "Цена"}, 3, "Прайс")
    again.delete("ЭлектроОпт")
    assert ProfileStore().get("ЭлектроОпт") is None


def test_profile_store_tolerates_bad_file(store_path):
    store_path.parent.mkdir(parents=True)
    store_path.write_text("{ это не json", encoding="utf-8")
    assert ProfileStore().names() == []
    store_path.write_text(json.dumps({
        "a": {"mapping": {"price": "Цена", "unknown": "X"}, "header_row": "2"},
        "b": "мусор",
    }), encoding="utf-8")
    store = ProfileStore()
    assert store.names() == ["a"]
    assert store.get("a") == SupplierProfile({PRICE: "Цена"}, None, None)


def test_apply_saved_profile(make_xlsx):
    path = make_xlsx("p.xlsx", [["Шапка"], ["Код", "Товар", "Опт", "Розница"], ["1", "A", "10", "15"]])
    source = SupplierSource.from_file(path)
    assert source.apply_saved({SKU: "Код", NAME: "Товар", PRICE: "Розница"}, 1, None)
    assert source.mapping[PRICE] == "Розница"
    # Профиль от другого файла не подходит — остаётся автоподбор
    assert not source.apply_saved({PRICE: "Нет такой колонки"}, 1, None)
    assert not source.apply_saved({PRICE: "Опт"}, 1, "Нет листа")
    assert source.mapping[PRICE] == "Розница"
    assert source.header_row == 1


def test_parse_tier():
    tier = cli.parse_tier("1000:15,5")
    assert (tier.min_price, tier.percent) == (1000, 15.5)
    with pytest.raises(Exception):
        cli.parse_tier("abc")


def test_cli_end_to_end(demo_files, tmp_path, store_path, capsys):
    output = tmp_path / "итог.xlsx"
    code = cli.main([*map(str, demo_files), "-o", str(output), "--markup", "20",
                     "--tier", "1000:10", "--round", "10"])
    assert code == 0
    printed = capsys.readouterr().out
    assert "Готово:" in printed
    assert "автоподбор" in printed
    sheet = load_workbook(output)["Итог"]
    sale_prices = [row[8].value for row in sheet.iter_rows(min_row=2) if row[8].value]
    assert sale_prices and all(price % 10 == 0 for price in sale_prices)


def test_cli_uses_profile(demo_files, tmp_path, store_path, capsys):
    ProfileStore().save("ЭлектроОпт", SupplierProfile({NAME: "Наименование товара", PRICE: "Цена, руб."}, 4))
    assert cli.main([str(demo_files[0]), "-o", str(tmp_path / "a.xlsx")]) == 0
    assert "ЭлектроОпт: строка заголовка 5, профиль" in capsys.readouterr().out


def test_cli_error(tmp_path, store_path, capsys):
    assert cli.main([str(tmp_path / "нет.xlsx")]) == 1
    assert "не найден" in capsys.readouterr().err
