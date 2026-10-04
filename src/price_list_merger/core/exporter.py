"""Выгрузка результата в Excel.

Листы итогового файла:
- «Итог» — по строке на товар: лучшая цена, поставщик, цена продажи
  и цены всех поставщиков (минимальная подсвечена зелёным);
- «Все предложения» — каждая строка каждого прайса и с чем она сопоставлена;
- «Не сопоставлено» — товары, которые нашлись только у одного поставщика;
- «Отброшено» — строки, которые не удалось разобрать, с причиной;
- «Сводка» — цифры и настройки, с которыми сделан свод.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .merging import price_column
from .pipeline import MergeReport
from .pricing import ROUNDING_LABELS

MONEY = "#,##0.00"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
MAX_COLUMN_WIDTH = 60


def export_report(report: MergeReport, path: str | Path) -> Path:
    """Сохраняет результат в .xlsx и возвращает путь к файлу."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    book = Workbook()
    _write_products(book.active, report)
    _write_offers(book.create_sheet("Все предложения"), report)
    _write_unmatched(book.create_sheet("Не сопоставлено"), report)
    _write_rejected(book.create_sheet("Отброшено"), report)
    _write_summary(book.create_sheet("Сводка"), report)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def _write_products(sheet: Worksheet, report: MergeReport) -> None:
    sheet.title = "Итог"
    suppliers = report.suppliers
    headers = [
        "№", "Артикул", "Наименование", "Ед.", "Поставщиков",
        "Лучшая цена", "Поставщик", "Наценка, %", "Цена продажи",
        *[f"Цена: {name}" for name in suppliers],
    ]
    sheet.append(headers)
    first_price_column = 10  # колонка J — цена первого поставщика
    for _, product in report.products.iterrows():
        row = [
            int(product["group"]) + 1, product["sku"], product["name"], product["unit"],
            int(product["offers"]), _value(product["best_price"]), product["best_supplier"],
            _value(product["markup"]), _value(product["sale_price"]),
            *[_value(product[price_column(name)]) for name in suppliers],
        ]
        sheet.append(row)
        excel_row = sheet.max_row
        # Подсвечиваем ячейку с минимальной ценой среди поставщиков
        if product["best_supplier"]:
            column = first_price_column + suppliers.index(product["best_supplier"])
            sheet.cell(excel_row, column).fill = BEST_FILL
        if pd.isna(product["best_price"]):
            sheet.cell(excel_row, 6).fill = WARN_FILL
    _format_table(sheet, money_columns=[6, 9, *range(first_price_column, first_price_column + len(suppliers))])


def _write_offers(sheet: Worksheet, report: MergeReport) -> None:
    sheet.append([
        "№ товара", "Поставщик", "Строка в файле", "Артикул", "Наименование",
        "Цена", "Остаток", "Ед.", "Сопоставлено по", "Похожесть, %", "Лучшая цена",
    ])
    offers = report.offers.sort_values(["group", "price"], kind="stable")
    for _, offer in offers.iterrows():
        sheet.append([
            int(offer["group"]) + 1, offer["supplier"], int(offer["row"]), offer["sku"],
            offer["name"], _value(offer["price"]), _value(offer["stock"]), offer["unit"],
            offer["match"], _value(offer["score"]) if offer["score"] < 100 else None,
            "да" if offer["best"] else "",
        ])
        if offer["best"]:
            for cell in sheet[sheet.max_row]:
                cell.fill = BEST_FILL
    _format_table(sheet, money_columns=[6])


def _write_unmatched(sheet: Worksheet, report: MergeReport) -> None:
    sheet.append(["№ товара", "Поставщик", "Артикул", "Наименование", "Цена"])
    products = report.products[report.products["offers"] == 1]
    for _, product in products.iterrows():
        supplier = product["best_supplier"] or _only_supplier(product, report.suppliers)
        sheet.append([
            int(product["group"]) + 1, supplier, product["sku"], product["name"],
            _value(product["best_price"]),
        ])
    _format_table(sheet, money_columns=[5])


def _write_rejected(sheet: Worksheet, report: MergeReport) -> None:
    sheet.append(["Поставщик", "Строка в файле", "Причина", "Содержимое строки"])
    for _, row in report.rejected.iterrows():
        sheet.append([row["supplier"], int(row["row"]), row["reason"], row["content"]])
    _format_table(sheet)


def _write_summary(sheet: Worksheet, report: MergeReport) -> None:
    stats = report.stats()
    pricing = report.pricing
    tiers = ", ".join(
        f"от {tier.min_price:g} ₽ — {tier.percent:g}%"
        for tier in sorted(pricing.tiers, key=lambda item: item.min_price)
    )
    rows = [
        ("Дата свода", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Поставщиков", stats["suppliers"]),
        ("Прочитано строк", stats["rows"]),
        ("Принято предложений", stats["offers"]),
        ("Отброшено строк", stats["rejected"]),
        ("Товаров в итоге", stats["products"]),
        ("Товаров у нескольких поставщиков", stats["in_several"]),
        ("Товаров, сопоставленных по похожему названию", stats["fuzzy"]),
        ("Наценка по умолчанию, %", pricing.percent),
        ("Ступени наценки", tiers or "нет"),
        ("Округление", ROUNDING_LABELS.get(pricing.rounding, pricing.rounding)),
    ]
    for supplier in report.suppliers:
        rows.append((f"Строк у «{supplier}»", report.rows_read.get(supplier, 0)))
    sheet.append(["Показатель", "Значение"])
    for row in rows:
        sheet.append(list(row))
    _format_table(sheet, autofilter=False)


def _only_supplier(product: pd.Series, suppliers: list[str]) -> str:
    """Поставщик товара, у которого нет лучшей цены (цена у всех пустая)."""
    for name in suppliers:
        if pd.notna(product[price_column(name)]):
            return name
    return ""


def _value(value: object) -> object:
    """NaN и None в Excel пишем как пустую ячейку."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def _format_table(sheet: Worksheet, money_columns: list[int] = (), autofilter: bool = True) -> None:
    """Шапка, закреплённая первая строка, фильтр, формат денег и ширина колонок."""
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    if autofilter and sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions
    for column in money_columns:
        for (cell,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cell.number_format = MONEY
    for index, cells in enumerate(sheet.iter_cols(), start=1):
        longest = max((len(str(cell.value)) for cell in cells if cell.value is not None), default=8)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), MAX_COLUMN_WIDTH)
